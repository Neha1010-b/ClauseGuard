"""
Document Loader — Phase 1.1
Extracts raw text + structural metadata from PDF / DOCX / TXT files.

Design principles:
- Single public function: load_document(path) -> LoadedDocument
- Format detected from file extension (never trust file contents alone)
- Raises clear exceptions on unsupported or empty files
- Never modifies the input file (raw data is sacred)
"""
from pathlib import Path
from dataclasses import dataclass, field
from typing import Optional, List

import fitz  # PyMuPDF
import docx


# ----------------------------------------------------------------
# Data model
# ----------------------------------------------------------------
@dataclass
class PageText:
    """Text from a single page (PDF) or the whole file (DOCX/TXT)."""
    page_number: int
    text: str
    char_start: int   # offset of this page's text in the full document string
    char_end: int


@dataclass
class LoadedDocument:
    """Result of loading a document — always the same shape regardless of format."""
    source_path: Path
    format: str                     # "pdf" | "docx" | "txt"
    full_text: str                  # concatenated, in document order
    pages: List[PageText]           # for PDFs; single-element for others
    metadata: dict = field(default_factory=dict)


# ----------------------------------------------------------------
# Format-specific extractors
# ----------------------------------------------------------------
def _load_pdf(path: Path) -> LoadedDocument:
    """Extract text page-by-page, tracking char offsets."""
    doc = fitz.open(path)
    pages: List[PageText] = []
    chunks: List[str] = []
    cursor = 0

    for i, page in enumerate(doc):
        text = page.get_text("text")
        if not text:
            continue
        chunks.append(text)
        pages.append(PageText(
            page_number=i + 1,
            text=text,
            char_start=cursor,
            char_end=cursor + len(text),
        ))
        cursor += len(text)

    metadata = {
        "page_count": doc.page_count,
        "pdf_metadata": dict(doc.metadata or {}),
    }
    doc.close()

    return LoadedDocument(
        source_path=path,
        format="pdf",
        full_text="".join(chunks),
        pages=pages,
        metadata=metadata,
    )


def _load_docx(path: Path) -> LoadedDocument:
    """Extract paragraph text from a .docx file."""
    d = docx.Document(str(path))
    paragraphs = [p.text for p in d.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)

    pages = [PageText(page_number=1, text=text, char_start=0, char_end=len(text))]
    metadata = {
        "paragraph_count": len(paragraphs),
        "core_properties": {
            "author": d.core_properties.author,
            "title": d.core_properties.title,
            "created": str(d.core_properties.created),
        },
    }
    return LoadedDocument(
        source_path=path,
        format="docx",
        full_text=text,
        pages=pages,
        metadata=metadata,
    )


def _load_txt(path: Path) -> LoadedDocument:
    """Read plain text with encoding fallback."""
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = path.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"Could not decode {path} with any supported encoding")

    pages = [PageText(page_number=1, text=text, char_start=0, char_end=len(text))]
    return LoadedDocument(
        source_path=path,
        format="txt",
        full_text=text,
        pages=pages,
        metadata={"encoding_used": enc},
    )


# ----------------------------------------------------------------
# Public API
# ----------------------------------------------------------------
_DISPATCH = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".txt": _load_txt,
}


def load_document(path: str | Path) -> LoadedDocument:
    """
    Load a document. Format is inferred from extension.
    Raises FileNotFoundError, ValueError (unsupported format), or RuntimeError (empty).
    """
    path = Path(path).resolve()
    if not path.exists():
        raise FileNotFoundError(f"No such file: {path}")

    ext = path.suffix.lower()
    if ext not in _DISPATCH:
        raise ValueError(
            f"Unsupported format '{ext}'. Supported: {list(_DISPATCH.keys())}"
        )

    doc = _DISPATCH[ext](path)

    if not doc.full_text.strip():
        raise RuntimeError(
            f"No extractable text in {path.name}. "
            "The file may be a scanned PDF requiring OCR."
        )
    return doc